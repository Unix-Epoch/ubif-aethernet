#!/usr/bin/env python3
"""
Provoke's sheet sync: downloads the store's lead texts (C.Boyd's "NewportNews
Lead sheet.docx") and the Price sheet (via ubif-workorders/fetch_price_sheet.py,
which also refreshes the audit's copy) through the signed-in Chrome profile,
parses both and pushes them to the Worker (POST /v1/sheets/<store>/<kind>).
Aethernet reads them from there, so staff need no SharePoint access and the
script needs no new permissions.

Every tab of the Price sheet is sent except Eagle of Va (Garrett, Oct 6): as raw rows, plus device
rows ({name, cells: {column: raw}}) for tabs laid out as a device-by-repair grid.

Only when a file changed: each run asks SharePoint for both files' version tags
(ETag, no download) and downloads, parses and pushes only the sheets whose tag
differs from the last push (sheets/.state.json). Each pushed sheet carries the
file's own last-modified time ("modified"), which Aethernet shows.

Key: the SHEETS_KEY Worker secret, read from the macOS Keychain
(service "aethernet-sheets"). Runs from launchd every 15 minutes; outside
8 AM-8 PM it does nothing.

Usage: ~/Scripts/ubif-daily-sales/.venv/bin/python3 sync_sheets.py [--force] [--dry-run] [--any-hour]
"""
import io
import json
import re
import subprocess
import sys
import urllib.request
import zipfile
from datetime import datetime
from pathlib import Path

sys.path.insert(0, str(Path.home() / "Scripts" / "ubif-workorders"))
import fetch_price_sheet as fps  # noqa: E402

STORE = 914
WORKER = "https://aethernet-waymark-dev.ubifaethernet.workers.dev"
APP_KEY = "ae_QFa2DKDy0DrlBd_A3bfZL3BtgM51idh6"  # public: ships inside Aethernet
SP = "https://ubreakifix-my.sharepoint.com/personal/"
FILES = {  # kind: (owner's OneDrive, SharePoint UniqueId)
    "leads": ("c_boyd_ubreakifix_com", "02F051BD-D255-4E00-BC98-F621D65F3284"),
    "prices": ("b_bauer_ubreakifix_com", fps.UNIQUE_ID),
}
HERE = Path(__file__).resolve().parent
LEADS_FILE = HERE / "lead-sheet.docx"
STATE = HERE / ".state.json"


def keychain():
    return subprocess.run(["security", "find-generic-password", "-s", "aethernet-sheets", "-w"],
                          capture_output=True, text=True, check=True).stdout.strip()


# ------------------------------------------------------------------ leads (.docx)
def paragraphs(docx):
    xml = zipfile.ZipFile(io.BytesIO(docx)).read("word/document.xml").decode("utf-8")
    out = []
    for p in re.findall(r"<w:p[ >].*?</w:p>", xml, re.S):
        t = "".join(re.findall(r"<w:t[^>]*>([^<]*)</w:t>", p))
        t = (t.replace("&amp;", "&").replace("&lt;", "<").replace("&gt;", ">").replace("&quot;", '"').replace("&apos;", "'"))
        t = re.sub(r"\s+", " ", t).strip()
        if t:
            out.append(t)
    return out


def is_message(t):
    return bool(re.match(r"(hi|hello|here at)\b", t, re.I)) or len(t) > 110


def parse_leads(docx):
    """[{group, title, texts}] in the doc's order. A title with no message under it
    before the next title is a group heading (INSURANCE LEADS, Misc Partner Leads, ...)."""
    items, group, title, texts = [], "", None, []

    def flush():
        if title is not None and texts:
            items.append({"group": group, "title": title.rstrip(" -:"), "texts": list(texts)})

    for t in paragraphs(docx):
        if is_message(t):
            texts.append(t)
            continue
        if title is not None and not texts:
            group = title.rstrip(" -:")  # the previous line was a heading
        flush()
        title, texts = t, []
    flush()
    return {"items": items}


# ------------------------------------------------------------------ prices (.xlsx)
PRICE_RE = re.compile(r"\d")
NOT_PRICE = re.compile(r"(n/?a|-+|.*self repair)$", re.I)  # a wide row's "no price" cell, not a repair name


def cell(v):
    return "" if v is None else re.sub(r"\s+", " ", str(v)).strip()


def is_header(c):
    """A row of column names: 3+ labels after the first cell and no prices (the console
    tab starts a new one per brand: "Sony Play Station | | HDMI / RETIMER | HDD 500GB ...")."""
    rest = [v for v in c[1:] if v]
    return len(rest) >= 3 and not any(re.match(r"\$?\s*\d", v) or NOT_PRICE.match(v) for v in rest)


def parse_prices(xlsx_path):
    """{tabs: [{name, raw: [[cells]], rows: [{name, cells: {column: raw text}}]}]}. Raw text is
    kept ("$190 (MS)", "N/A", "Apple self repair"); Aethernet reads the number from it.
    The Google tab also has a long layout (device | repair | price, sometimes two side by side)."""
    from openpyxl import load_workbook
    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    tabs = []
    for ws in wb.worksheets:
        if ws.title.strip().lower().startswith("eagle"):  # contract customer tab: ignored (Garrett, Oct 6)
            continue
        rows, header, raw = {}, None, []
        for r in ws.iter_rows(values_only=True):
            c = [cell(v) for v in r]
            while c and not c[-1]:
                c.pop()
            if c:
                raw.append(c)
            if not c or not c[0]:
                continue
            if header is None or is_header(c):
                header = c
                continue
            # long layout: name | repair | price (repeated in threes)
            if len(header) > 1 and header[1] and len(c) >= 2 and c[1] and not PRICE_RE.search(c[1]) and not NOT_PRICE.match(c[1]) and (len(c) < 3 or PRICE_RE.search(c[2])):
                c += [""] * (-len(c) % 3)
                for i in range(0, len(c), 3):
                    name, repair, price = c[i], c[i + 1], c[i + 2]
                    if name and repair and price:
                        rows.setdefault(name, {"name": name, "cells": {}})["cells"][repair] = price
                continue
            row = rows.setdefault(c[0], {"name": c[0], "cells": {}})
            for i, v in enumerate(c[1:], 1):
                col = header[i] if i < len(header) and header[i] else ("Notes" if i == 1 else "")
                if col and v:
                    row["cells"][col] = v
        tabs.append({"name": ws.title, "raw": raw, "rows": list(rows.values())})
    return {"tabs": tabs}


# ------------------------------------------------------------------ download + push
def fetch_changed(seen, force):
    """{kind: (bytes, etag, modified)} for the files whose version tag changed (all with force)."""
    from playwright.sync_api import sync_playwright
    out = {}
    with sync_playwright() as p:
        ctx = p.chromium.launch_persistent_context(user_data_dir=str(fps.PROFILE_DIR), channel="chrome", headless=True)
        page = ctx.pages[0] if ctx.pages else ctx.new_page()
        page.goto("https://ubreakifix-my.sharepoint.com/", timeout=30000, wait_until="domcontentloaded")
        page.wait_for_timeout(1500)
        try:
            for kind, (owner, uid) in FILES.items():
                r = ctx.request.get(f"{SP}{owner}/_api/web/GetFileById('{uid}')?$select=ETag,TimeLastModified",
                                    headers={"accept": "application/json;odata=nometadata"})
                if r.status != 200:
                    raise RuntimeError(f"{kind}: version check failed, HTTP {r.status}")
                info = r.json()
                if not force and seen.get(kind) == info["ETag"]:
                    continue
                d = ctx.request.get(f"{SP}{owner}/_layouts/15/download.aspx?UniqueId={uid}")
                body = d.body()
                if d.status != 200 or not body.startswith(b"PK"):
                    raise RuntimeError(f"{kind}: download failed, HTTP {d.status}, {len(body)} bytes")
                out[kind] = (body, info["ETag"], info["TimeLastModified"])
        finally:
            ctx.close()
    return out


def push(kind, data, key):
    req = urllib.request.Request(f"{WORKER}/v1/sheets/{STORE}/{kind}", method="POST",
                                 data=json.dumps({"data": data}).encode(),
                                 headers={"content-type": "application/json", "x-aethernet-key": APP_KEY,
                                          "x-aethernet-sheets": key, "user-agent": "aethernet-sheets"})
    urllib.request.urlopen(req, timeout=20).read()


def main():
    force, dry = "--force" in sys.argv, "--dry-run" in sys.argv
    if not force and "--any-hour" not in sys.argv and not (8 <= datetime.now().hour < 20):
        return 0
    seen = json.loads(STATE.read_text()) if STATE.exists() else {}
    lock = fps.acquire_profile_lock()
    if lock is None:
        print("Timed out waiting for the Chrome profile lock.", file=sys.stderr)
        return 2
    try:
        changed = fetch_changed(seen, force)
    finally:
        lock.close()
    if not changed:
        print("no changes")
        return 0
    data = {}
    if "leads" in changed:
        LEADS_FILE.write_bytes(changed["leads"][0])
        data["leads"] = parse_leads(changed["leads"][0])
    if "prices" in changed:
        fps.OUT_FILE.write_bytes(changed["prices"][0])  # also keeps the audit's copy current
        fps.export_text()
        data["prices"] = parse_prices(fps.OUT_FILE)
    for kind in data:
        data[kind]["modified"] = changed[kind][2]
    print("changed: " + ", ".join(f"{k} (edited {changed[k][2]})" for k in data))
    if dry:
        return 0
    key = keychain()
    for kind, d in data.items():
        push(kind, d, key)
        seen[kind] = changed[kind][1]
    STATE.write_text(json.dumps(seen))
    print("pushed")
    return 0


if __name__ == "__main__":
    sys.exit(main())
